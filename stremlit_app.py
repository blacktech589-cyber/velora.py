"""Binomo için CSV tabanlı araştırma ve Excel sinyal uygulaması.

Bu yazılım yatırım tavsiyesi veya otomatik işlem aracı değildir. Binomo'dan
dışa aktarılan mum verisini analiz eder; platform hesabına bağlanmaz.
"""
from __future__ import annotations

import io
import logging
from datetime import datetime
from pathlib import Path

import joblib
import numpy as np
import pandas as pd
import streamlit as st
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from sklearn.ensemble import ExtraTreesClassifier, RandomForestClassifier
from sklearn.impute import SimpleImputer
from sklearn.metrics import accuracy_score, precision_score, recall_score
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import StandardScaler

try:
    import tensorflow as tf
    from tensorflow.keras import Model, Sequential
    from tensorflow.keras.callbacks import EarlyStopping
    from tensorflow.keras.layers import (
        GRU, LSTM, Conv1D, Dense, Dropout, GlobalAveragePooling1D, Input
    )
    TF_AVAILABLE = True
except ImportError:
    TF_AVAILABLE = False


APP_DIR = Path(__file__).resolve().parent
MODEL_DIR = APP_DIR / "models"
OUTPUT_FILE = APP_DIR / "binomo_sinyalleri.xlsx"
LOG_FILE = APP_DIR / "hata_log.txt"
MODEL_DIR.mkdir(exist_ok=True)
logging.basicConfig(
    filename=LOG_FILE,
    level=logging.INFO,
    format="%(asctime)s %(levelname)s %(message)s",
    encoding="utf-8",
)

FEATURES = [
    # Getiriler ve trend
    "return_1", "return_3", "return_5", "return_10",
    "sma_ratio_5", "sma_ratio_10", "sma_ratio_20", "sma_ratio_50",
    "ema_ratio_5", "ema_ratio_9", "ema_ratio_12", "ema_ratio_15",
    "ema_ratio_21", "ema_ratio_26", "ema_ratio_50",
    "ema_15_slope", "ema_15_26_spread",
    # Momentum
    "rsi_7", "rsi_14", "rsi_21", "stoch_k", "stoch_d",
    "williams_r", "roc_10", "cci_20",
    # Trend gücü
    "macd", "macd_signal", "macd_histogram", "adx_14", "plus_di", "minus_di",
    # Oynaklık ve fiyat konumu
    "volatility_10", "volatility_20", "atr_14", "natr_14",
    "bb_position", "bb_width", "bb_percent_b", "range_ratio",
    # Hacim
    "volume_change", "obv_change", "mfi_14", "vwap_ratio",
]


def normalize_columns(df: pd.DataFrame) -> pd.DataFrame:
    aliases = {
        "timestamp": "time", "datetime": "time", "date": "time",
        "tarih": "time", "zaman": "time", "açılış": "open", "acilis": "open",
        "yüksek": "high", "yuksek": "high", "düşük": "low", "dusuk": "low",
        "kapanış": "close", "kapanis": "close", "hacim": "volume",
    }
    out = df.copy()
    out.columns = [aliases.get(str(c).strip().lower(), str(c).strip().lower()) for c in out.columns]
    required = {"time", "open", "high", "low", "close"}
    missing = required.difference(out.columns)
    if missing:
        raise ValueError("Eksik sütun(lar): " + ", ".join(sorted(missing)))
    if "volume" not in out:
        out["volume"] = 0.0
    out["time"] = pd.to_datetime(out["time"], errors="coerce", utc=True)
    for col in ["open", "high", "low", "close", "volume"]:
        out[col] = pd.to_numeric(out[col], errors="coerce")
    out = out.dropna(subset=["time", "open", "high", "low", "close"])
    out = out.sort_values("time").drop_duplicates("time").reset_index(drop=True)
    if len(out) < 150:
        raise ValueError("Eğitim için en az 150 geçerli mum gerekir.")
    return out


def rsi(close: pd.Series, period: int = 14) -> pd.Series:
    delta = close.diff()
    gain = delta.clip(lower=0).ewm(alpha=1 / period, adjust=False).mean()
    loss = (-delta.clip(upper=0)).ewm(alpha=1 / period, adjust=False).mean()
    rs = gain / loss.replace(0, np.nan)
    return (100 - 100 / (1 + rs)).fillna(50) / 100


def true_range(frame: pd.DataFrame) -> pd.Series:
    previous_close = frame["close"].shift(1)
    return pd.concat([
        frame["high"] - frame["low"],
        (frame["high"] - previous_close).abs(),
        (frame["low"] - previous_close).abs(),
    ], axis=1).max(axis=1)


def directional_index(frame: pd.DataFrame, period: int = 14):
    up = frame["high"].diff()
    down = -frame["low"].diff()
    plus_dm = pd.Series(np.where((up > down) & (up > 0), up, 0.0), index=frame.index)
    minus_dm = pd.Series(np.where((down > up) & (down > 0), down, 0.0), index=frame.index)
    atr = true_range(frame).ewm(alpha=1 / period, adjust=False).mean()
    plus_di = 100 * plus_dm.ewm(alpha=1 / period, adjust=False).mean() / atr.replace(0, np.nan)
    minus_di = 100 * minus_dm.ewm(alpha=1 / period, adjust=False).mean() / atr.replace(0, np.nan)
    dx = 100 * (plus_di - minus_di).abs() / (plus_di + minus_di).replace(0, np.nan)
    adx = dx.ewm(alpha=1 / period, adjust=False).mean()
    return adx / 100, plus_di / 100, minus_di / 100


def money_flow_index(frame: pd.DataFrame, period: int = 14) -> pd.Series:
    typical = (frame["high"] + frame["low"] + frame["close"]) / 3
    flow = typical * frame["volume"]
    positive = flow.where(typical.diff() > 0, 0.0).rolling(period).sum()
    negative = flow.where(typical.diff() < 0, 0.0).rolling(period).sum()
    ratio = positive / negative.replace(0, np.nan)
    return (100 - 100 / (1 + ratio)).fillna(50) / 100


def add_features(df: pd.DataFrame, horizon: int) -> pd.DataFrame:
    x = df.copy()
    close = x["close"]
    for n in [1, 3, 5, 10]:
        x[f"return_{n}"] = close.pct_change(n)
    for n in [5, 10, 20, 50]:
        x[f"sma_ratio_{n}"] = close / close.rolling(n).mean() - 1
    emas = {}
    for n in [5, 9, 12, 15, 21, 26, 50]:
        emas[n] = close.ewm(span=n, adjust=False).mean()
        x[f"ema_ratio_{n}"] = close / close.ewm(span=n, adjust=False).mean() - 1
    x["ema_15_slope"] = emas[15].pct_change(3)
    x["ema_15_26_spread"] = emas[15] / emas[26] - 1
    returns = close.pct_change()
    x["volatility_10"] = returns.rolling(10).std()
    x["volatility_20"] = returns.rolling(20).std()
    x["rsi_7"] = rsi(close, 7)
    x["rsi_14"] = rsi(close)
    x["rsi_21"] = rsi(close, 21)
    ema12, ema26 = emas[12], emas[26]
    x["macd"] = (ema12 - ema26) / close
    x["macd_signal"] = (ema12 - ema26).ewm(span=9, adjust=False).mean() / close
    x["macd_histogram"] = x["macd"] - x["macd_signal"]
    mid = close.rolling(20).mean()
    deviation = close.rolling(20).std()
    upper, lower = mid + 2 * deviation, mid - 2 * deviation
    band_range = (upper - lower).replace(0, np.nan)
    x["bb_position"] = (close - lower) / band_range
    x["bb_percent_b"] = x["bb_position"]
    x["bb_width"] = band_range / mid.replace(0, np.nan)
    low14 = x["low"].rolling(14).min()
    high14 = x["high"].rolling(14).max()
    x["stoch_k"] = (close - low14) / (high14 - low14).replace(0, np.nan)
    x["stoch_d"] = x["stoch_k"].rolling(3).mean()
    x["williams_r"] = (close - high14) / (high14 - low14).replace(0, np.nan)
    x["roc_10"] = close.pct_change(10)
    typical = (x["high"] + x["low"] + close) / 3
    mean_deviation = typical.rolling(20).apply(
        lambda values: np.mean(np.abs(values - values.mean())), raw=True
    )
    x["cci_20"] = (typical - typical.rolling(20).mean()) / (0.015 * mean_deviation).replace(0, np.nan)
    atr = true_range(x).ewm(alpha=1 / 14, adjust=False).mean()
    x["atr_14"] = atr / close.replace(0, np.nan)
    x["natr_14"] = x["atr_14"]
    x["adx_14"], x["plus_di"], x["minus_di"] = directional_index(x)
    x["range_ratio"] = (x["high"] - x["low"]) / close.replace(0, np.nan)
    x["volume_change"] = x["volume"].pct_change().replace([np.inf, -np.inf], np.nan)
    direction = np.sign(close.diff()).fillna(0)
    obv = (direction * x["volume"]).cumsum()
    x["obv_change"] = obv.pct_change(5).replace([np.inf, -np.inf], np.nan)
    x["mfi_14"] = money_flow_index(x)
    cumulative_volume = x["volume"].replace(0, np.nan).cumsum()
    vwap = (typical * x["volume"]).cumsum() / cumulative_volume
    x["vwap_ratio"] = close / vwap.replace(0, np.nan) - 1
    future_return = close.shift(-horizon) / close - 1
    x["target"] = np.where(future_return.notna(), (future_return > 0).astype(int), np.nan)
    return x.replace([np.inf, -np.inf], np.nan)


def sequences(frame: pd.DataFrame, lookback: int):
    values = frame[FEATURES].to_numpy(dtype=np.float32)
    labels = frame["target"].to_numpy()
    xs, ys, rows = [], [], []
    for end in range(lookback - 1, len(frame)):
        if np.isnan(labels[end]):
            continue
        window = values[end - lookback + 1:end + 1]
        if np.isnan(window).all(axis=0).any():
            continue
        med = np.nanmedian(window, axis=0)
        window = np.where(np.isnan(window), med, window)
        xs.append(window)
        ys.append(int(labels[end]))
        rows.append(end)
    return np.asarray(xs), np.asarray(ys), rows


def build_dl_models(shape):
    early = EarlyStopping(monitor="val_loss", patience=5, restore_best_weights=True)
    lstm = Sequential([
        Input(shape=shape), LSTM(48), Dropout(.2), Dense(24, activation="relu"),
        Dense(1, activation="sigmoid")
    ], name="LSTM")
    gru = Sequential([
        Input(shape=shape), GRU(48), Dropout(.2), Dense(24, activation="relu"),
        Dense(1, activation="sigmoid")
    ], name="GRU")
    inp = Input(shape=shape)
    z = Conv1D(48, 3, padding="causal", activation="relu")(inp)
    z = Conv1D(24, 3, padding="causal", activation="relu")(z)
    z = GlobalAveragePooling1D()(z)
    cnn = Model(inp, Dense(1, activation="sigmoid")(z), name="CNN1D")
    for model in (lstm, gru, cnn):
        model.compile(optimizer="adam", loss="binary_crossentropy", metrics=["accuracy"])
    return [lstm, gru, cnn], early


def train_and_predict(frame, lookback, epochs):
    X, y, row_ids = sequences(frame, lookback)
    if len(X) < 100 or len(np.unique(y)) < 2:
        raise ValueError("Yeterli ve iki sınıfı da içeren eğitim örneği bulunamadı.")
    split = int(len(X) * .8)
    if split < 50 or len(X) - split < 20:
        raise ValueError("Zaman bazlı test bölümü için daha fazla mum gerekiyor.")
    X_train, X_test, y_train, y_test = X[:split], X[split:], y[:split], y[split:]
    flat_train = X_train.reshape(len(X_train), -1)
    flat_test = X_test.reshape(len(X_test), -1)
    latest_seq = X[-1:]
    probabilities, names = [], []

    classical = [
        ("RandomForest", RandomForestClassifier(
            n_estimators=300, min_samples_leaf=3, class_weight="balanced",
            random_state=42, n_jobs=-1)),
        ("ExtraTrees", ExtraTreesClassifier(
            n_estimators=300, min_samples_leaf=3, class_weight="balanced",
            random_state=42, n_jobs=-1)),
    ]
    for name, estimator in classical:
        pipe = Pipeline([("imputer", SimpleImputer()), ("scale", StandardScaler()), ("model", estimator)])
        pipe.fit(flat_train, y_train)
        probabilities.append(pipe.predict_proba(flat_test)[:, 1])
        names.append(name)
        joblib.dump(pipe, MODEL_DIR / f"{name}.joblib")

    if TF_AVAILABLE:
        dl_models, early = build_dl_models(X_train.shape[1:])
        for model in dl_models:
            model.fit(
                X_train, y_train, validation_split=.2, epochs=epochs, batch_size=32,
                callbacks=[early], verbose=0, shuffle=False,
            )
            probabilities.append(model.predict(X_test, verbose=0).ravel())
            names.append(model.name)
            model.save(MODEL_DIR / f"{model.name}.keras")

    ensemble = np.mean(probabilities, axis=0)
    pred = (ensemble >= .5).astype(int)
    metrics = {
        "Doğruluk": accuracy_score(y_test, pred),
        "Precision": precision_score(y_test, pred, zero_division=0),
        "Recall": recall_score(y_test, pred, zero_division=0),
        "Test örneği": len(y_test),
    }

    latest_probs = []
    for name, estimator in classical:
        pipe = joblib.load(MODEL_DIR / f"{name}.joblib")
        latest_probs.append(float(pipe.predict_proba(latest_seq.reshape(1, -1))[0, 1]))
    if TF_AVAILABLE:
        for model in dl_models:
            latest_probs.append(float(model.predict(latest_seq, verbose=0)[0, 0]))
    probability = float(np.mean(latest_probs))
    return probability, metrics, names


def append_excel(record: dict) -> bytes:
    new = pd.DataFrame([record])
    if OUTPUT_FILE.exists():
        old = pd.read_excel(OUTPUT_FILE, sheet_name="Sinyaller")
        data = pd.concat([old, new], ignore_index=True).tail(5000)
    else:
        data = new
    with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
        data.to_excel(writer, sheet_name="Sinyaller", index=False)
        ws = writer.sheets["Sinyaller"]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center")
        for col in ws.columns:
            letter = col[0].column_letter
            ws.column_dimensions[letter].width = min(32, max(12, max(len(str(c.value or "")) for c in col) + 2))
        if ws.max_row > 1:
            ws.conditional_formatting.add(
                f"D2:D{ws.max_row}",
                ColorScaleRule(start_type="min", start_color="F8696B",
                               mid_type="percentile", mid_value=50, mid_color="FFEB84",
                               end_type="max", end_color="63BE7B"),
            )
        pd.DataFrame([{
            "Not": "Araştırma amaçlıdır; yatırım tavsiyesi veya kazanç garantisi değildir.",
            "Veri": "Kullanıcının yüklediği geçmiş OHLC mum verisi",
            "Doğrulama": "Son %20 veri, zaman sırası korunarak test edilir.",
        }]).to_excel(writer, sheet_name="Açıklama", index=False)
    return OUTPUT_FILE.read_bytes()


def latest_indicators(frame: pd.DataFrame) -> dict:
    row = frame.iloc[-1]
    return {
        "EMA 15": round(float(row["close"] / (1 + row["ema_ratio_15"])), 6),
        "EMA15 eğimi": round(float(row["ema_15_slope"]), 6),
        "RSI 14": round(float(row["rsi_14"] * 100), 2),
        "Bollinger %B": round(float(row["bb_percent_b"]), 4),
        "Bollinger genişliği": round(float(row["bb_width"]), 6),
        "MACD": round(float(row["macd"]), 6),
        "MACD histogram": round(float(row["macd_histogram"]), 6),
        "ATR %": round(float(row["atr_14"] * 100), 4),
        "ADX": round(float(row["adx_14"] * 100), 2),
        "Stochastic %K": round(float(row["stoch_k"] * 100), 2),
        "Williams %R": round(float(row["williams_r"] * 100), 2),
        "CCI 20": round(float(row["cci_20"]), 2),
        "MFI 14": round(float(row["mfi_14"] * 100), 2),
    }


st.set_page_config(page_title="Binomo DL Araştırma", layout="wide")
st.title("Binomo CSV Derin Öğrenme Araştırması")
st.warning("Araştırma amaçlıdır. Otomatik işlem yapmaz; yatırım tavsiyesi veya kazanç garantisi değildir.")
uploaded = st.file_uploader("Binomo mum verisi CSV", type=["csv"])
c1, c2, c3 = st.columns(3)
asset = c1.text_input("Varlık", "EUR/USD")
horizon = c2.number_input("Tahmin ufku (mum)", 1, 20, 1)
lookback = c3.number_input("Model penceresi (mum)", 20, 120, 40)
epochs = st.slider("DL epoch", 5, 100, 25)

if uploaded and st.button("Eğit, test et ve Excel'e kaydet", type="primary"):
    try:
        raw = pd.read_csv(uploaded, sep=None, engine="python")
        data = normalize_columns(raw)
        frame = add_features(data, int(horizon))
        with st.spinner("Modeller eğitiliyor..."):
            probability, metrics, model_names = train_and_predict(frame, int(lookback), epochs)
        indicators = latest_indicators(frame)
        signal = "YUKARI" if probability >= .5 else "AŞAĞI"
        confidence = probability if signal == "YUKARI" else 1 - probability
        record = {
            "Zaman": datetime.now().astimezone().isoformat(timespec="seconds"),
            "Varlık": asset,
            "Sinyal": signal,
            "Model olasılığı": round(probability, 4),
            "Güven": round(confidence, 4),
            "Tahmin ufku (mum)": horizon,
            "Model sayısı": len(model_names),
            "Modeller": ", ".join(model_names),
            "Test doğruluğu": round(metrics["Doğruluk"], 4),
            "Test precision": round(metrics["Precision"], 4),
            "Test recall": round(metrics["Recall"], 4),
            "Mum sayısı": len(data),
            **indicators,
        }
        excel_bytes = append_excel(record)
        st.success(f"Sonuç: {signal} — model güveni %{confidence * 100:.1f}")
        st.json(metrics)
        st.subheader("Son mum teknik göstergeleri")
        st.dataframe(pd.DataFrame([indicators]), use_container_width=True, hide_index=True)
        st.caption("Kullanılan modeller: " + ", ".join(model_names))
        st.download_button(
            "Excel'i indir", excel_bytes, OUTPUT_FILE.name,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    except Exception as exc:
        logging.exception("Analiz başarısız")
        st.error(f"İşlem tamamlanamadı: {exc}")
